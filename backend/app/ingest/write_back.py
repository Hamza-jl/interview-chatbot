"""Write confirmed answers back into the file that was uploaded.

The original file is opened and modified, never regenerated. Styles, headers,
numbering, pre-printed instructions and the author's formatting all survive,
because the only thing that changes is the contents of the cells the reader
found. That is the whole point: the deliverable has to look like the document
the client sent, not like something a program produced.

Every answer is addressed by the `target` its slot recorded at discovery time,
so nothing is searched for a second time. A document that shifted between upload
and download would otherwise be filled from stale coordinates - which is why the
uploaded bytes are what get filled, not a fresh copy of anything.

Word reuses `app.pca.docx_filler`, which already handles the hard parts: keeping
an "N.B." instruction under a written answer, and cloning a template row so new
grid rows carry the document's own borders and fonts.
"""
from __future__ import annotations

import io
import logging
from typing import Any, Dict, List, Optional

from .slots import DocumentPlan, Slot

logger = logging.getLogger("pca.ingest.write_back")


class WriteBackError(Exception):
    """Raised with a message meant for the person who uploaded the file."""


def _value(payload: Dict[str, Any]) -> str:
    return str(payload.get("value") or "").strip()


def _rows(payload: Dict[str, Any]) -> List[Dict[str, str]]:
    rows = payload.get("rows")
    return [r for r in rows if isinstance(r, dict) and any(r.values())] if rows else []


# --------------------------------------------------------------------------- #
# Word
# --------------------------------------------------------------------------- #
def _fill_docx(source: bytes, plan: DocumentPlan, answers: Dict[str, Dict[str, Any]]) -> bytes:
    from docx import Document

    from app.pca.docx_filler import fill_grid, set_cell

    document = Document(io.BytesIO(source))
    tables = document.tables

    for slot in plan.slots:
        payload = answers.get(slot.id)
        if not payload:
            continue
        index = slot.target["table"] - 1
        if not 0 <= index < len(tables):
            logger.warning("slot %s points at table %s, absent", slot.id, index + 1)
            continue
        table = tables[index]

        if slot.kind == "grid":
            rows = _rows(payload)
            if rows:
                fill_grid(
                    table,
                    [c.id for c in _columns(slot)],
                    rows,
                    header_rows=slot.target.get("header_rows", 1),
                )
            continue

        text = _value(payload)
        if not text:
            continue
        row, column = slot.target["row"], slot.target["col"]
        if row >= len(table.rows) or column >= len(table.rows[row].cells):
            logger.warning("slot %s points outside table %s", slot.id, index + 1)
            continue
        set_cell(table.rows[row].cells[column], text)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class _Column:
    """`fill_grid` expects objects with an `id`; slots carry plain labels."""

    __slots__ = ("id",)

    def __init__(self, identifier: str) -> None:
        self.id = identifier


def _columns(slot: Slot) -> List[_Column]:
    return [_Column(name) for name in slot.columns]


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _fill_xlsx(source: bytes, plan: DocumentPlan, answers: Dict[str, Dict[str, Any]]) -> bytes:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(source))

    for slot in plan.slots:
        payload = answers.get(slot.id)
        if not payload:
            continue
        name = slot.target["sheet"]
        if name not in workbook.sheetnames:
            logger.warning("slot %s names sheet %r, absent", slot.id, name)
            continue
        sheet = workbook[name]

        if slot.kind == "grid":
            rows = _rows(payload)
            first = slot.target.get("header_row", 1) + 1
            column0 = slot.target.get("first_column", 1)
            for offset, row in enumerate(rows):
                for index, name_ in enumerate(slot.columns):
                    sheet.cell(
                        row=first + offset,
                        column=column0 + index,
                        value=row.get(name_, ""),
                    )
            continue

        text = _value(payload)
        if text:
            sheet.cell(row=slot.target["row"], column=slot.target["col"], value=text)

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# PowerPoint
# --------------------------------------------------------------------------- #
def _fill_pptx(source: bytes, plan: DocumentPlan, answers: Dict[str, Dict[str, Any]]) -> bytes:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(source))
    slides = list(deck.slides)

    for slot in plan.slots:
        payload = answers.get(slot.id)
        if not payload:
            continue
        number = slot.target["slide"] - 1
        if not 0 <= number < len(slides):
            continue
        shapes = list(slides[number].shapes)
        order = slot.target["shape"]
        if not 0 <= order < len(shapes):
            continue
        shape = shapes[order]

        if slot.kind == "grid" and shape.has_table:
            table = shape.table
            rows = _rows(payload)
            # A deck's table cannot grow without rebuilding the shape, so extra
            # answers are reported rather than silently dropped.
            available = len(table.rows) - 1
            for offset, row in enumerate(rows[:available]):
                for index, name in enumerate(slot.columns):
                    if index < len(table.columns):
                        table.cell(offset + 1, index).text = row.get(name, "")
            if len(rows) > available:
                logger.warning(
                    "slot %s: %s rows given, %s available in the slide table",
                    slot.id, len(rows), available,
                )
            continue

        text = _value(payload)
        if text and shape.has_text_frame:
            shape.text_frame.text = text

    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _fill_pdf(source: bytes, plan: DocumentPlan, answers: Dict[str, Dict[str, Any]]) -> bytes:
    import fitz

    document = fitz.open(stream=source, filetype="pdf")
    if not document.is_form_pdf:
        document.close()
        raise WriteBackError(
            "Ce PDF n'a pas de champs de formulaire : il n'y a pas d'emplacement "
            "où écrire. Fournissez le document source (Word, Excel)."
        )

    by_name: Dict[str, str] = {}
    for slot in plan.slots:
        payload = answers.get(slot.id)
        if payload and slot.target.get("field_name"):
            text = _value(payload)
            if text:
                by_name[slot.target["field_name"]] = text

    for page in document:
        for widget in page.widgets() or []:
            value = by_name.get(widget.field_name)
            if value is None:
                continue
            widget.field_value = value
            widget.update()

    buffer = io.BytesIO()
    document.save(buffer)
    document.close()
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
_WRITERS = {"docx": _fill_docx, "xlsx": _fill_xlsx, "pptx": _fill_pptx, "pdf": _fill_pdf}


def fill(source: bytes, plan: DocumentPlan, answers: Dict[str, Dict[str, Any]]) -> bytes:
    """Return the uploaded file with every confirmed answer written into it.

    `answers` maps slot id to the same payload shape the interview already
    stores: ``{"value": str}`` or ``{"rows": [{column: value}]}``. A slot with no
    answer is left exactly as it was - a blank in the original stays a blank,
    which is what makes an unfinished document honest.
    """
    writer = _WRITERS.get(plan.kind)
    if writer is None:
        raise WriteBackError(f"Format non pris en charge en écriture : {plan.kind}")
    return writer(source, plan, answers)
