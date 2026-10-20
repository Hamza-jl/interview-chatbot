"""Upload a document, answer its questions, get it back filled.

The round trip is the only assertion that matters here. Every intermediate step
can look correct while the answer still lands in the wrong cell, so these tests
open the *produced* file and read the cell the answer was supposed to reach.
"""
from __future__ import annotations

import io

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from app.core.config import settings
from tests.conftest import API, authenticate, confirm, say


@pytest.fixture()
def headers(client):
    return authenticate(client, "client@example.com")


def _questionnaire(tmp_path):
    """Two blanks and one empty table - the shape these templates really have."""
    document = Document()
    document.add_heading("Fiche d'identité de l'entité", level=1)
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Nom du responsable"
    table.rows[1].cells[0].text = "Effectif"

    document.add_heading("Applications utilisées", level=1)
    grid = document.add_table(rows=3, cols=2)
    grid.rows[0].cells[0].text = "Application"
    grid.rows[0].cells[1].text = "Criticité"

    path = tmp_path / "questionnaire.docx"
    document.save(path)
    return path


def _upload(client, headers, path):
    with open(path, "rb") as handle:
        res = client.post(
            f"{API}/documents", headers=headers,
            files={"file": (path.name, handle.read())},
        )
    assert res.status_code == 201, res.text
    return res.json()


# --------------------------------------------------------------------------- #
# Reading
# --------------------------------------------------------------------------- #
def test_an_upload_reports_what_it_found(client, headers, tmp_path):
    body = _upload(client, headers, _questionnaire(tmp_path))

    assert body["kind"] == "docx"
    assert body["slots"] == 3, body["preview"]
    assert body["grids"] == 1
    assert "Nom du responsable" in [p["label"] for p in body["preview"]]


def test_a_text_file_is_now_read_rather_than_refused(client, headers, tmp_path):
    """Avant l'etape distante, un .txt etait refuse : rien a remplir dedans.

    Il est desormais lu. Un fichier non remplissable produit quand meme un
    entretien valable ; ses reponses sortiront dans un nouveau document plutot
    que dans l'original.
    """
    path = tmp_path / "cadrage.txt"
    path.write_text("Perimetre a documenter : correspondant PCA, effectif, "
                    "activites critiques.", encoding="utf-8")

    with open(path, "rb") as handle:
        res = client.post(f"{API}/documents", headers=headers,
                          files={"file": ("cadrage.txt", handle.read())})

    assert res.status_code == 201, res.text
    assert res.json()["kind"] == "txt"


def test_a_file_larger_than_the_json_limit_still_gets_through(client, headers, tmp_path):
    """Observé en vrai : un vrai modèle BIA de 583 Ko refusé en 413.

    Le plafond global des requêtes est à 256 Ko, ce qui convient à une réponse
    d'entretien et pas du tout à un fichier. Sans exception explicite le refus
    tombe avant la route, donc avant tout message parlant du document - et
    l'écran, lui, annonce 12 Mo.
    """
    path = tmp_path / "gros.txt"
    path.write_text("Perimetre a documenter.\n" * 30_000, encoding="utf-8")
    assert path.stat().st_size > settings.MAX_BODY_BYTES

    with open(path, "rb") as handle:
        res = client.post(f"{API}/documents", headers=headers,
                          files={"file": ("gros.txt", handle.read())})

    assert res.status_code == 201, res.text


def test_a_file_beyond_the_upload_limit_is_refused(client, headers, monkeypatch):
    """L'exception relève le plafond, elle ne le supprime pas."""
    monkeypatch.setattr(settings, "MAX_UPLOAD_BYTES", 64 * 1024, raising=False)

    res = client.post(f"{API}/documents", headers=headers,
                      files={"file": ("enorme.txt", b"x" * (128 * 1024))})

    assert res.status_code == 413
    assert "volumineuse" in res.json()["detail"].lower()


def test_a_format_nobody_can_read_is_still_refused(client, headers, tmp_path):
    """Un .doc de 1997 n'est pas lisible : le refus doit dire quoi envoyer."""
    path = tmp_path / "vieux.doc"
    # En-tete OLE2 : la signature des formats Office 97-2003.
    path.write_bytes(bytes.fromhex("d0cf11e0a1b11ae1") + bytes(64))

    with open(path, "rb") as handle:
        res = client.post(f"{API}/documents", headers=headers,
                          files={"file": ("vieux.doc", handle.read())})

    assert res.status_code == 415
    assert "97-2003" in res.json()["detail"]


def test_a_document_belongs_to_the_account_that_uploaded_it(client, headers, tmp_path):
    body = _upload(client, headers, _questionnaire(tmp_path))
    other = authenticate(client, "staff@example.org")

    res = client.get(f"{API}/documents/{body['id']}", headers=other)
    assert res.status_code == 404


# --------------------------------------------------------------------------- #
# The interview
# --------------------------------------------------------------------------- #
def test_questions_are_written_for_every_blank(client, headers, tmp_path):
    body = _upload(client, headers, _questionnaire(tmp_path))

    res = client.post(f"{API}/documents/{body['id']}/questions", headers=headers)
    assert res.status_code == 200, res.text
    questions = res.json()["questions"]

    assert len(questions) == 3
    assert all(q["question"].strip() for q in questions)
    grid = next(q for q in questions if q["kind"] == "grid")
    assert grid["columns"] == ["Application", "Criticité"]


def test_the_interview_runs_on_the_ordinary_chat_endpoints(client, headers, tmp_path):
    """The whole point of registering under doc:<id>: nothing else changes."""
    body = _upload(client, headers, _questionnaire(tmp_path))
    client.post(f"{API}/documents/{body['id']}/questions", headers=headers)

    res = client.post(f"{API}/documents/{body['id']}/session", headers=headers)
    assert res.status_code == 201, res.text
    detail = res.json()
    session_id = detail["state"]["id"]
    assert detail["state"]["total"] == 3

    turn = say(client, headers, session_id, "Mme Leila Ben Salah")
    assert turn["pending"] is not None
    confirm(client, headers, session_id, turn["pending"])

    state = client.get(f"{API}/sessions/{session_id}", headers=headers).json()["state"]
    assert state["answered"] == 1


def test_a_second_start_resumes_rather_than_duplicating(client, headers, tmp_path):
    body = _upload(client, headers, _questionnaire(tmp_path))
    client.post(f"{API}/documents/{body['id']}/questions", headers=headers)

    first = client.post(f"{API}/documents/{body['id']}/session", headers=headers).json()
    second = client.post(f"{API}/documents/{body['id']}/session", headers=headers).json()

    assert first["state"]["id"] == second["state"]["id"]


# --------------------------------------------------------------------------- #
# The deliverable
# --------------------------------------------------------------------------- #
def _run_interview(client, headers, document_id, answers):
    client.post(f"{API}/documents/{document_id}/questions", headers=headers)
    detail = client.post(f"{API}/documents/{document_id}/session", headers=headers).json()
    session_id = detail["state"]["id"]
    for reply in answers:
        turn = say(client, headers, session_id, reply)
        if turn.get("pending"):
            confirm(client, headers, session_id, turn["pending"])
    return session_id


def _download(client, headers, produced):
    res = client.get(f"{API}/exports/{produced['download_token']}", headers=headers)
    assert res.status_code == 200, res.text
    return res


def test_the_answers_land_in_the_cells_they_were_read_from(client, headers, tmp_path):
    body = _upload(client, headers, _questionnaire(tmp_path))
    _run_interview(client, headers, body["id"], ["Mme Leila Ben Salah", "42 personnes"])

    produced = client.post(f"{API}/documents/{body['id']}/fill", headers=headers)
    assert produced.status_code == 200, produced.text
    produced = produced.json()
    assert produced["filename"] == "rempli_questionnaire.docx"
    assert produced["answers"] == 2

    filled = Document(io.BytesIO(_download(client, headers, produced).content))
    table = filled.tables[0]
    assert table.rows[0].cells[1].text.strip() == "Mme Leila Ben Salah"
    assert "42" in table.rows[1].cells[1].text
    # The label column is never touched: it is the document's own wording.
    assert table.rows[0].cells[0].text.strip() == "Nom du responsable"


def test_an_unanswered_blank_stays_blank(client, headers, tmp_path):
    body = _upload(client, headers, _questionnaire(tmp_path))
    _run_interview(client, headers, body["id"], ["Mme Leila Ben Salah"])

    produced = client.post(f"{API}/documents/{body['id']}/fill", headers=headers).json()
    filled = Document(io.BytesIO(_download(client, headers, produced).content))

    assert produced["blanks"] == 2
    assert filled.tables[0].rows[1].cells[1].text.strip() == ""


def test_filling_before_any_answer_is_refused(client, headers, tmp_path):
    body = _upload(client, headers, _questionnaire(tmp_path))
    client.post(f"{API}/documents/{body['id']}/questions", headers=headers)
    client.post(f"{API}/documents/{body['id']}/session", headers=headers)

    res = client.post(f"{API}/documents/{body['id']}/fill", headers=headers)
    assert res.status_code == 409
    assert "identique" in res.json()["detail"]


def test_an_excel_comes_back_as_excel(client, headers, tmp_path):
    """Serving a workbook as Word makes Excel refuse to open its own file."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Contacts"
    sheet["A1"], sheet["A2"] = "Responsable PCA", "Suppléant"
    path = tmp_path / "contacts.xlsx"
    book.save(path)

    body = _upload(client, headers, path)
    assert body["kind"] == "xlsx"
    _run_interview(client, headers, body["id"], ["Karim Trabelsi"])

    produced = client.post(f"{API}/documents/{body['id']}/fill", headers=headers).json()
    response = _download(client, headers, produced)

    assert "spreadsheetml" in response.headers["content-type"]
    filled = load_workbook(io.BytesIO(response.content))
    assert filled["Contacts"]["B1"].value == "Karim Trabelsi"


def test_deleting_a_document_takes_its_questions_with_it(client, headers, tmp_path):
    from app.ingest import plan as ingest_plan

    body = _upload(client, headers, _questionnaire(tmp_path))
    client.post(f"{API}/documents/{body['id']}/questions", headers=headers)
    assert ingest_plan.is_registered(body["id"])

    assert client.delete(f"{API}/documents/{body['id']}", headers=headers).status_code == 204
    assert not ingest_plan.is_registered(body["id"])
    assert client.get(f"{API}/documents/{body['id']}", headers=headers).status_code == 404


def test_a_discovered_blank_is_a_cell_not_a_free_answer(client, headers, tmp_path):
    """The extraction prompt only says "copy it verbatim" for a field.

    Mapped to "open", the model reformulates: "42 collaborateurs" came back as
    "nettoye pour 42 collaborateurs", and that is what was written into the
    client's own document.
    """
    body = _upload(client, headers, _questionnaire(tmp_path))
    client.post(f"{API}/documents/{body['id']}/questions", headers=headers)
    detail = client.post(f"{API}/documents/{body['id']}/session", headers=headers).json()

    kinds = {q["question_id"]: q["kind"] for q in
             client.get(f"{API}/sessions/{detail['state']['id']}/answers",
                        headers=headers).json()}

    assert "open" not in kinds.values(), kinds
    assert sorted(set(kinds.values())) == ["field", "grid"]


# --------------------------------------------------------------------------- #
# The briefing
# --------------------------------------------------------------------------- #
BRIEFING = (
    "Division Réassurance, 4 personnes, rattachée à la Direction Technique. "
    "Entretien de mise à jour : la fiche 2025 existe, seuls les processus ont changé."
)


def test_the_briefing_is_saved_and_read_back(client, headers, tmp_path):
    body = _upload(client, headers, _questionnaire(tmp_path))

    res = client.put(f"{API}/documents/{body['id']}/context", headers=headers,
                     json={"context": BRIEFING})
    assert res.status_code == 200, res.text

    detail = client.get(f"{API}/documents/{body['id']}", headers=headers).json()
    assert detail["context"] == BRIEFING
    assert detail["has_context"] is True


def test_the_briefing_is_sealed_at_rest(client, headers, tmp_path):
    """It is the freest-typed field in the application, so the likeliest to
    carry a name, a headcount or a weakness somebody would rather not leak."""
    from sqlalchemy import select
    from app.db.models import UploadedDocument
    from app.db.session import SessionLocal

    body = _upload(client, headers, _questionnaire(tmp_path))
    client.put(f"{API}/documents/{body['id']}/context", headers=headers,
               json={"context": BRIEFING})

    with SessionLocal() as db:
        stored = db.execute(
            select(UploadedDocument.context_enc)
            .where(UploadedDocument.id == body["id"])
        ).scalar_one()

    assert stored, "something must be stored"
    assert "Réassurance" not in stored
    assert "4 personnes" not in stored


def test_the_briefing_reaches_the_question_writer(client, headers, tmp_path):
    """The whole point: the model must be told what the document cannot say."""
    from app.ingest import questions as ingest_questions

    body = _upload(client, headers, _questionnaire(tmp_path))
    client.put(f"{API}/documents/{body['id']}/context", headers=headers,
               json={"context": BRIEFING})

    seen = []

    class Recording:
        label = "stub"

        def structured(self, *, system, messages, **_):
            seen.append(messages[0]["content"])
            return {"question": "Quel est l'effectif de la division ?",
                    "aide": "", "exemple": "4"}

    from app.ai import llm

    saved = llm.active_backend
    llm.active_backend = lambda: Recording()
    try:
        res = client.post(f"{API}/documents/{body['id']}/questions", headers=headers)
    finally:
        llm.active_backend = saved

    assert res.status_code == 200, res.text
    assert seen, "the writer was never called"
    assert all("CONTEXTE DE LA MISSION" in prompt for prompt in seen)
    assert all("Division Réassurance" in prompt for prompt in seen)


def test_an_empty_briefing_adds_nothing_to_the_prompt(tmp_path):
    from app.ingest import questions as ingest_questions
    from app.ingest.slots import Slot

    slot = Slot(id="s", kind="cell", label="Effectif", target={}, context=["Identité"])
    assert "CONTEXTE" not in ingest_questions.prompt_for(slot, "")
    assert "CONTEXTE" in ingest_questions.prompt_for(slot, BRIEFING)


def test_the_briefing_is_flagged_as_not_being_an_answer():
    """A briefing saying "4 personnes" must not become the answer to "Effectif".

    The interview prompt labels it as the consultant's own note, so the model
    has no licence to read it back as something the interviewee said.
    """
    from app.ai import engine
    from app.pca.blueprint import get_plan

    question = get_plan("entite")[0]
    block = engine.build_context_block(
        question,
        structure_name="Division Réassurance",
        template_label="entite",
        position="question 1 sur 10",
        existing=None,
        followups=0,
        background=BRIEFING,
    )

    assert "CONTEXTE DE LA MISSION" in block
    assert "pas une reponse de la personne interrogee" in block


# --------------------------------------------------------------------------- #
# Upgrading a database that is already in service
# --------------------------------------------------------------------------- #
def test_the_new_columns_are_added_to_an_existing_database(tmp_path):
    """A deployed server has survey_sessions already; create_all never alters it.

    Left unregistered, the first interview after the upgrade fails on
    "no column named document_id" - on the server, not here.
    """
    import sqlalchemy
    from app.db import session as db_session

    path = tmp_path / "old.db"
    engine = sqlalchemy.create_engine(f"sqlite+pysqlite:///{path.as_posix()}")

    # The schema as it stood before this feature.
    with engine.begin() as connection:
        connection.execute(sqlalchemy.text(
            "CREATE TABLE survey_sessions (id VARCHAR(32) PRIMARY KEY, "
            "template_kind VARCHAR(16))"
        ))
        connection.execute(sqlalchemy.text(
            "INSERT INTO survey_sessions (id, template_kind) VALUES ('s1', 'entite')"
        ))

    saved = db_session.engine
    db_session.engine = engine
    try:
        db_session._apply_pending_columns()
    finally:
        db_session.engine = saved

    columns = {
        c["name"]
        for c in sqlalchemy.inspect(engine).get_columns("survey_sessions")
    }
    assert "document_id" in columns

    # And the interview that was already open is still there.
    with engine.begin() as connection:
        rows = connection.execute(
            sqlalchemy.text("SELECT id, template_kind FROM survey_sessions")
        ).all()
    assert rows == [("s1", "entite")]
    engine.dispose()


def test_a_document_template_kind_is_longer_than_the_old_column(tmp_path):
    """The reason the widening exists, stated as a number.

    Postgres enforces VARCHAR length; "doc:" + a 32-character id does not fit
    in the 16 characters the column was declared with.
    """
    from app.ingest import plan as ingest_plan

    kind = ingest_plan.template_kind("0" * 32)
    assert len(kind) == 36
    assert len(kind) > 16
