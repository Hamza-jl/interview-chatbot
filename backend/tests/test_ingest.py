"""Discovering the blanks in an uploaded document.

Structure discovery is deterministic by design, so it is testable by design.
These tests build documents whose correct answer is known, then assert the
reader finds exactly that - because the failure mode being guarded against is
not "no slots found", it is "a slot pointing at the wrong cell", which is
invisible until an answer lands in the wrong place in a signed document.
"""
from __future__ import annotations

import pytest
from docx import Document
from openpyxl import Workbook

from app.ingest import questions, read_docx, read_xlsx, reader
from app.ingest.slots import Slot, is_blank
from tests.conftest import needs_templates


# --------------------------------------------------------------------------- #
# Telling a blank from an answer
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    "text",
    ["", "   ", "…", "....", "___", "-----", "[à compléter]", "(nom)", "N/A", "néant", "xxx"],
)
def test_furniture_counts_as_blank(text):
    assert is_blank(text)


@pytest.mark.parametrize("text", ["Mme Leila Ben Salah", "12", "Comptabilité", "0"])
def test_a_real_answer_is_not_blank(text):
    assert not is_blank(text)


# --------------------------------------------------------------------------- #
# Word
# --------------------------------------------------------------------------- #
def _docx_with_keyvalue(tmp_path, header: bool):
    document = Document()
    document.add_heading("1. Fiche de suivi", level=1)
    rows = ["Nom du responsable", "Direction"]
    table = document.add_table(rows=len(rows) + (1 if header else 0), cols=2)
    offset = 0
    if header:
        table.rows[0].cells[0].text = "Question"
        table.rows[0].cells[1].text = "Réponse"
        offset = 1
    for i, label in enumerate(rows):
        table.rows[i + offset].cells[0].text = label
    path = tmp_path / f"kv-{'header' if header else 'plain'}.docx"
    document.save(path)
    return path


@pytest.mark.parametrize("header", [False, True])
def test_a_label_beside_an_empty_cell_becomes_a_question(tmp_path, header):
    """'Question | Réponse' is how these templates are really written."""
    plan = read_docx.read(str(_docx_with_keyvalue(tmp_path, header)))

    labels = [s.label for s in plan.cells]
    assert labels == ["Nom du responsable", "Direction"], labels
    assert all(s.target["mode"] == "cell" for s in plan.cells)
    assert all(s.target["col"] == 1 for s in plan.cells), "answers go in column 2"


def test_the_header_row_is_never_asked_about(tmp_path):
    plan = read_docx.read(str(_docx_with_keyvalue(tmp_path, header=True)))
    assert "Question" not in [s.label for s in plan.cells]


def test_a_header_over_empty_rows_becomes_one_table_question(tmp_path):
    document = Document()
    document.add_heading("5. Applications", level=1)
    table = document.add_table(rows=4, cols=3)
    for i, name in enumerate(["Domaine", "Application", "Criticité"]):
        table.rows[0].cells[i].text = name
    path = tmp_path / "grid.docx"
    document.save(path)

    plan = read_docx.read(str(path))
    assert len(plan.grids) == 1 and not plan.cells
    grid = plan.grids[0]
    assert grid.columns == ["Domaine", "Application", "Criticité"]
    assert grid.context[-1] == "5. Applications", "the heading gives the question meaning"


def test_a_filled_table_is_not_mistaken_for_a_form(tmp_path):
    """Data already in the document must not be turned into questions."""
    document = Document()
    table = document.add_table(rows=3, cols=2)
    for row, (a, b) in enumerate([("Domaine", "Responsable"),
                                  ("Crédit", "Mme Ben Salah"),
                                  ("Risques", "M. Trabelsi")]):
        table.rows[row].cells[0].text = a
        table.rows[row].cells[1].text = b
    path = tmp_path / "filled.docx"
    document.save(path)

    plan = read_docx.read(str(path))
    assert not plan.slots
    assert plan.notes, "a document with nothing to fill says so"


def test_an_author_instruction_is_a_hint_not_an_answer(tmp_path):
    """'N.B. : template envoyé séparément' is guidance left in the answer cell."""
    document = Document()
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Question"
    table.rows[0].cells[1].text = "Réponse"
    table.rows[1].cells[0].text = "Liste des documents critiques"
    table.rows[1].cells[1].text = "N.B. : template spécifique envoyé par mail"
    path = tmp_path / "hint.docx"
    document.save(path)

    plan = read_docx.read(str(path))
    assert len(plan.cells) == 1, "the row is still a question"
    assert "template" in (plan.cells[0].existing or ""), "the note is kept as a hint"


@needs_templates
def test_the_real_templates_yield_the_sections_they_contain():
    """The strongest available evidence that discovery works on real documents."""
    plan = read_docx.read("templates/etat_des_lieux_entite.docx")

    labels = [s.label for s in plan.cells]
    assert "Nom du Responsable" in labels
    assert "Entité" in labels

    sections = {s.context[-1] for s in plan.grids if s.context}
    assert any("Activités" in s for s in sections)
    assert plan.grids, "the repeating tables must be found"
    assert all(len(g.columns) >= 2 for g in plan.grids)


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def test_a_spreadsheet_form_is_read_even_when_column_b_is_empty(tmp_path):
    """openpyxl reports max_column == 1 for a blank answer column."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Fiche"
    for row, label in enumerate(["Nom du responsable", "Direction"], start=1):
        sheet.cell(row=row, column=1, value=label)
    path = tmp_path / "form.xlsx"
    book.save(path)

    plan = read_xlsx.read(str(path))
    assert [s.label for s in plan.cells] == ["Nom du responsable", "Direction"]
    assert all(s.target["col"] == 2 for s in plan.cells)


def test_a_sheet_with_only_a_header_is_still_a_table(tmp_path):
    """The commonest shape of a blank grid has no rows under the header at all."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Applications"
    for column, name in enumerate(["Domaine", "Application", "Criticité"], start=1):
        sheet.cell(row=1, column=column, value=name)
    path = tmp_path / "grid.xlsx"
    book.save(path)

    plan = read_xlsx.read(str(path))
    assert len(plan.grids) == 1
    assert plan.grids[0].columns == ["Domaine", "Application", "Criticité"]


# --------------------------------------------------------------------------- #
# Choosing a reader
# --------------------------------------------------------------------------- #
def test_the_format_is_detected_from_content_not_the_extension(tmp_path):
    """An .xlsx renamed .docx must not reach the Word reader."""
    book = Workbook()
    book.active["A1"] = "Nom"
    disguised = tmp_path / "piege.docx"
    book.save(disguised)

    assert reader.detect(str(disguised)) == "xlsx"


@pytest.mark.parametrize("name,fragment", [
    ("vieux.doc", "97-2003"),
    ("donnees.csv", "CSV"),
    ("notes.txt", "texte"),
])
def test_an_unsupported_format_is_refused_with_a_reason(tmp_path, name, fragment):
    path = tmp_path / name
    path.write_bytes(b"pas une archive")
    with pytest.raises(reader.UnsupportedDocument) as refused:
        reader.detect(str(path), name)
    assert fragment.lower() in str(refused.value).lower()


# --------------------------------------------------------------------------- #
# Phrasing a question
# --------------------------------------------------------------------------- #
def _grid_slot() -> Slot:
    return Slot(
        id="t1.grid", kind="grid", label="Applications",
        columns=["Domaine", "Application", "Criticité"],
        context=["5. Applications"], target={},
    )


def test_a_table_example_always_matches_the_column_count():
    """The row parser downstream needs this exact, so it is not left to a model.

    Measured on the real templates, qwen2.5:3b returned the right number of
    cells 37% of the time and qwen3:4b 14%. Padding here makes it 100%.
    """
    slot = _grid_slot()
    for cells in ([], ["Crédit"], ["a", "b", "c"], ["a", "b", "c", "d", "e"]):
        row = questions.example_row(slot, cells)
        assert len(row.split(" | ")) == len(slot.columns), cells


def test_a_pipe_inside_a_cell_cannot_forge_a_column():
    row = questions.example_row(_grid_slot(), ["Crédit | Risque", "Delta", "V"])
    assert len(row.split(" | ")) == 3


def test_the_fallback_question_works_with_no_model_at_all():
    """Inference being down must not make an upload unusable."""
    generated = questions.fallback(_grid_slot())
    assert "Applications" in generated.question
    assert generated.engine == "deterministe"
    assert len(generated.example.split(" | ")) == 3


def test_the_prompt_names_the_columns_in_order():
    prompt = questions.prompt_for(_grid_slot())
    assert "Domaine | Application | Criticité" in prompt
    assert "5. Applications" in prompt


def test_the_parser_tolerates_prose_around_the_json():
    parsed = questions.parse(
        'Voici :\n```json\n{"question": "Quel domaine ?", "aide": "a", "cellules": ["x"]}\n```'
    )
    assert parsed and parsed["question"] == "Quel domaine ?"
    assert parsed["cellules"] == ["x"]


def test_an_answer_with_no_question_is_rejected():
    assert questions.parse('{"aide": "a", "exemple": "b"}') is None
    assert questions.parse("pas du json") is None
